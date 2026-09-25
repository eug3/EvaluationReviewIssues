#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从《华信电子底稿编制指引（征求意见稿）》的操作类底稿中抽取各科目《评估步骤及复核表》，
按科目归并去重，输出 questions/步骤复核表.yaml。

依赖：xlrd（.xls）、openpyxl（.xlsx）、pyyaml
用法：python3 scripts/extract-step-tables.py
"""
import collections
import os
import re
import sys

import xlrd
import openpyxl
import yaml

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(
    REPO,
    "华信电子底稿编制指引（征求意见稿）",
    "附件2：评【2025】-000 电子底稿",
    "2操作类底稿",
)
OUT = os.path.join(REPO, "questions", "步骤复核表.yaml")

STEP_RE = re.compile(r"^\s*\d{1,2}\s*[、.．]\s*(\S.{3,})")
INDEX_RE = re.compile(r"索引号[:：]?\s*(\S+)")
# 文件名前缀里的索引号，如 c3-5应收账款（C类项目）.xls → C3-5
FILE_INDEX_RE = re.compile(r"^([cC]\d(?:[-–]\d+)*)")
# 步骤表末尾的填表说明/免责声明，不是操作步骤，剔除
FOOTER_RE = re.compile(r"(仅供参考|转移相应责任|供评估人员现场清查|复核人提出追加程序|不因使用参考格式)")

# 科目别名归一：不同模板/文件夹的叫法统一到一个规范名
ALIASES = {
    "现金": "货币资金-现金",
    "货币资金—现金": "货币资金-现金",
    "货币资金—库存现金": "货币资金-现金",
    "银行存款": "货币资金-银行存款",
    "货币资金—银行存款": "货币资金-银行存款",
    "其他货币资金": "货币资金-其他货币资金",
    "货币资金—其他货币资金": "货币资金-其他货币资金",
    "库存商品": "产成品（库存商品）",
    "机器设备": "固定资产-机器设备",
    "固定资产—设备": "固定资产-机器设备",
    "固定资产－设备": "固定资产-机器设备",
    "固定资产－工程物资": "工程物资",
    "在建工程-土建": "在建工程—土建工程",
    "在建工程—土建": "在建工程—土建工程",
    "在建工程-安装": "在建工程—设备安装工程",
    "在建工程—设备安装": "在建工程—设备安装工程",
    "土地使用权": "无形资产-土地使用权",
    "无形资产—土地使用权": "无形资产-土地使用权",
    "无形资产-土地使用权": "无形资产-土地使用权",
    "无形资产—矿业权": "无形资产-矿业权",
    "无形资产—其他": "无形资产-其他",
    "建筑物": "固定资产-房屋建筑物/构筑物",
    "房屋建筑物": "固定资产-房屋建筑物/构筑物",
    "固定资产—房屋建筑物/构筑物": "固定资产-房屋建筑物/构筑物",
    "应付股利": "应付股利（应付利润）",
    "应收股利": "应收股利（应收利润）",
}


def clean_subject(title: str, fallback: str) -> str:
    s = INDEX_RE.sub("", title)
    s = re.sub(r"(评估)?步骤(及)?复核表", "", s)
    s = re.sub(r"\s+", " ", s).strip(" -—/")
    s = s or fallback
    return ALIASES.get(s, s)


def rows_of(path: str, sheet: str):
    if path.lower().endswith(".xls"):
        sh = xlrd.open_workbook(path).sheet_by_name(sheet)
        return [[str(sh.cell_value(r, c)).strip() for c in range(sh.ncols)] for r in range(sh.nrows)]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = [[("" if c is None else str(c).strip()) for c in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def main():
    records = []
    for dp, _dn, fn in os.walk(SRC):
        for f in sorted(fn):
            if f.startswith("~$"):
                continue
            low = f.lower()
            path = os.path.join(dp, f)
            if not low.endswith((".xls", ".xlsx", ".xlsm")):
                continue
            try:
                if low.endswith(".xls"):
                    names = xlrd.open_workbook(path).sheet_names()
                else:
                    wb = openpyxl.load_workbook(path, read_only=True)
                    names = wb.sheetnames
                    wb.close()
            except Exception as exc:  # 打不开的文件跳过，不中断
                print(f"SKIP {os.path.basename(path)}: {exc}", file=sys.stderr)
                continue
            for sheet in names:
                if "步骤" not in sheet:
                    continue
                try:
                    rows = rows_of(path, sheet)
                except Exception as exc:
                    print(f"SHEET SKIP {os.path.basename(path)}/{sheet}: {exc}", file=sys.stderr)
                    continue
                title, index, steps = "", "", []
                for cells in rows:
                    joined = " ".join(c for c in cells if c)
                    if not title and "复核表" in joined and len(joined) < 40:
                        title = joined
                    if not index:
                        m = INDEX_RE.search(joined)
                        if m:
                            index = m.group(1)
                    for c in cells:
                        m = STEP_RE.match(c)
                        if m:
                            step = re.sub(r"\s+", " ", m.group(1)).strip()
                            if len(step) >= 4 and not FOOTER_RE.search(step) and step not in steps:
                                steps.append(step)
                if steps:
                    file_index = FILE_INDEX_RE.match(os.path.basename(path))
                    records.append({
                        "subject": clean_subject(title, os.path.basename(dp)),
                        "index": file_index.group(1).upper() if file_index else index,
                        "file": os.path.relpath(path, REPO),
                        "steps": steps,
                    })

    by_subject = collections.OrderedDict()
    for rec in records:
        entry = by_subject.setdefault(rec["subject"], {"indexes": collections.Counter(), "files": [], "variants": {}})
        if rec["index"]:
            entry["indexes"][rec["index"]] += 1
        entry["files"].append(rec["file"])
        entry["variants"].setdefault(tuple(rec["steps"]), {"steps": rec["steps"], "count": 0})
        entry["variants"][tuple(rec["steps"])]["count"] += 1

    subjects = []
    for subject, entry in by_subject.items():
        # 同一科目存在多套步骤时取最全的一套
        best = max(entry["variants"].values(), key=lambda v: len(v["steps"]))
        # 索引号取该科目所有来源中出现次数最多的一个（旧模板可能带过期索引）
        index = entry["indexes"].most_common(1)[0][0] if entry["indexes"] else ""
        subjects.append({
            "subject": subject,
            "index": index,
            "step_count": len(best["steps"]),
            "variant_count": len(entry["variants"]),
            "sources": sorted(set(entry["files"])),
            "steps": best["steps"],
        })
    subjects.sort(key=lambda s: (s["index"] or "z", s["subject"]))

    doc = {
        "version": 1,
        "title": "各科目评估步骤及复核表（问题底稿）",
        "source": "华信电子底稿编制指引（征求意见稿）/附件2 操作类底稿",
        "extracted_by": "scripts/extract-step-tables.py",
        "note": "每条 step 对应底稿《步骤及复核表》中的一项操作步骤与要求（原表逐项勾选“是/不适用”）。",
        "summary": {
            "step_sheets": len(records),
            "subjects": len(subjects),
            "steps_total": sum(s["step_count"] for s in subjects),
        },
        "subjects": subjects,
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as fh:
        yaml.safe_dump(doc, fh, allow_unicode=True, sort_keys=False, width=120)
    print(f"step_sheets={len(records)} subjects={len(subjects)} steps={doc['summary']['steps_total']} -> {os.path.relpath(OUT, REPO)}")


if __name__ == "__main__":
    main()
